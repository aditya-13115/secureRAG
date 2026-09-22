from __future__ import annotations

import csv
import re
from dataclasses import dataclass
from pathlib import Path

import fitz  # PyMuPDF
from docx import Document as DocxDocument
from openpyxl import load_workbook


# ============================================================
# LOADED BLOCK
# ============================================================


@dataclass(slots=True)
class LoadedBlock:
    """
    A logically retrievable piece of a source document.

    The loader is responsible ONLY for extracting content and
    its source location.

    It does not know about:
        - users
        - permissions
        - embeddings
        - vector databases
        - LLMs
    """

    text: str

    # PDF
    page_number: int | None = None

    # DOCX / Markdown / TXT
    section: str | None = None

    # XLSX
    sheet_name: str | None = None
    row_start: int | None = None
    row_end: int | None = None

    # CSV / TXT / Markdown
    line_start: int | None = None
    line_end: int | None = None


# ============================================================
# TEXT CLEANING
# ============================================================


def clean_text(text: str) -> str:
    """
    Normalize extracted text without destroying meaningful
    paragraph boundaries.
    """

    text = text.replace("\x00", "")

    # Normalize Windows/Mac line endings.
    text = text.replace("\r\n", "\n")
    text = text.replace("\r", "\n")

    # Remove repeated spaces/tabs.
    text = re.sub(r"[ \t]+", " ", text)

    # Avoid huge numbers of blank lines.
    text = re.sub(r"\n{3,}", "\n\n", text)

    return text.strip()


# ============================================================
# PDF
# ============================================================


def load_pdf(path: Path) -> list[LoadedBlock]:
    """
    Extract one block per PDF page.

    Keeping the page boundary is important for citations.
    """

    blocks: list[LoadedBlock] = []

    with fitz.open(path) as pdf:

        for page_index, page in enumerate(pdf):

            text = clean_text(
                page.get_text("text")
            )

            if not text:
                continue

            blocks.append(
                LoadedBlock(
                    text=text,
                    page_number=page_index + 1,
                )
            )

    return blocks


# ============================================================
# DOCX
# ============================================================


def load_docx(path: Path) -> list[LoadedBlock]:
    """
    Extract DOCX paragraphs while preserving heading context.

    python-docx does not reliably expose rendered PDF-style
    page numbers, so section/heading information becomes the
    primary citation location.
    """

    document = DocxDocument(path)

    blocks: list[LoadedBlock] = []

    current_section: str | None = None

    paragraph_buffer: list[str] = []
    buffer_start = 0

    def flush_buffer(end_index: int) -> None:
        nonlocal paragraph_buffer
        nonlocal buffer_start

        if not paragraph_buffer:
            return

        text = clean_text(
            "\n\n".join(paragraph_buffer)
        )

        if text:
            blocks.append(
                LoadedBlock(
                    text=text,
                    section=current_section,
                    line_start=buffer_start + 1,
                    line_end=end_index + 1,
                )
            )

        paragraph_buffer = []

    for index, paragraph in enumerate(
        document.paragraphs
    ):

        text = clean_text(
            paragraph.text
        )

        if not text:
            continue

        style_name = (
            paragraph.style.name
            if paragraph.style
            else ""
        )

        # Heading 1 / Heading 2 / Heading 3
        # become citation sections.
        if style_name.startswith("Heading"):

            flush_buffer(index - 1)

            current_section = text

            continue

        if not paragraph_buffer:
            buffer_start = index

        paragraph_buffer.append(text)

    flush_buffer(
        len(document.paragraphs) - 1
    )

    return blocks


# ============================================================
# XLSX
# ============================================================


def load_xlsx(path: Path) -> list[LoadedBlock]:
    """
    Extract spreadsheet content while preserving:

        sheet name
        start row
        end row

    We group rows so one huge spreadsheet does not become a
    single enormous embedding.
    """

    workbook = load_workbook(
        filename=path,
        read_only=True,
        data_only=True,
    )

    blocks: list[LoadedBlock] = []

    rows_per_block = 25

    try:

        for worksheet in workbook.worksheets:

            current_rows: list[str] = []
            current_start: int | None = None
            current_end: int | None = None

            for row_number, row in enumerate(
                worksheet.iter_rows(
                    values_only=True
                ),
                start=1,
            ):

                values = []

                for value in row:
                    if value is None:
                        values.append("")
                    else:
                        values.append(
                            str(value).strip()
                        )

                # Skip completely empty rows.
                if not any(values):
                    continue

                row_text = " | ".join(values)

                if current_start is None:
                    current_start = row_number

                current_rows.append(
                    row_text
                )

                current_end = row_number

                if len(current_rows) >= rows_per_block:

                    text = clean_text(
                        "\n".join(current_rows)
                    )

                    blocks.append(
                        LoadedBlock(
                            text=text,
                            sheet_name=worksheet.title,
                            row_start=current_start,
                            row_end=current_end,
                        )
                    )

                    current_rows = []
                    current_start = None
                    current_end = None

            # Flush remaining rows.
            if current_rows:

                text = clean_text(
                    "\n".join(current_rows)
                )

                blocks.append(
                    LoadedBlock(
                        text=text,
                        sheet_name=worksheet.title,
                        row_start=current_start,
                        row_end=current_end,
                    )
                )

    finally:
        workbook.close()

    return blocks


# ============================================================
# CSV
# ============================================================


def load_csv(path: Path) -> list[LoadedBlock]:
    """
    Extract CSV rows while preserving line numbers.

    CSV is treated similarly to a spreadsheet.
    """

    blocks: list[LoadedBlock] = []

    rows_per_block = 25

    current_rows: list[str] = []
    current_start: int | None = None
    current_end: int | None = None

    with path.open(
        "r",
        encoding="utf-8-sig",
        newline="",
    ) as file:

        reader = csv.reader(file)

        for line_number, row in enumerate(
            reader,
            start=1,
        ):

            values = [
                str(value).strip()
                for value in row
            ]

            if not any(values):
                continue

            row_text = " | ".join(values)

            if current_start is None:
                current_start = line_number

            current_rows.append(
                row_text
            )

            current_end = line_number

            if len(current_rows) >= rows_per_block:

                blocks.append(
                    LoadedBlock(
                        text=clean_text(
                            "\n".join(current_rows)
                        ),
                        line_start=current_start,
                        line_end=current_end,
                    )
                )

                current_rows = []
                current_start = None
                current_end = None

    if current_rows:

        blocks.append(
            LoadedBlock(
                text=clean_text(
                    "\n".join(current_rows)
                ),
                line_start=current_start,
                line_end=current_end,
            )
        )

    return blocks


# ============================================================
# TXT / MARKDOWN
# ============================================================


def load_text(path: Path) -> list[LoadedBlock]:
    """
    Load TXT / Markdown.

    Markdown headings become section metadata.
    Plain TXT gets line ranges.
    """

    content = path.read_text(
        encoding="utf-8"
    )

    lines = content.splitlines()

    blocks: list[LoadedBlock] = []

    current_section: str | None = None

    paragraph_lines: list[str] = []
    paragraph_start: int | None = None

    def flush_paragraph(end_line: int) -> None:
        nonlocal paragraph_lines
        nonlocal paragraph_start

        if not paragraph_lines:
            return

        text = clean_text(
            "\n".join(paragraph_lines)
        )

        if text:

            blocks.append(
                LoadedBlock(
                    text=text,
                    section=current_section,
                    line_start=paragraph_start,
                    line_end=end_line,
                )
            )

        paragraph_lines = []
        paragraph_start = None

    for line_number, raw_line in enumerate(
        lines,
        start=1,
    ):

        line = raw_line.strip()

        # Markdown heading.
        heading_match = re.match(
            r"^#{1,6}\s+(.+)$",
            line,
        )

        if heading_match:

            flush_paragraph(
                line_number - 1
            )

            current_section = (
                heading_match.group(1).strip()
            )

            continue

        # Blank line = paragraph boundary.
        if not line:

            if paragraph_lines:
                flush_paragraph(
                    line_number - 1
                )

            continue

        if paragraph_start is None:
            paragraph_start = line_number

        paragraph_lines.append(line)

    flush_paragraph(
        len(lines)
    )

    return blocks


# ============================================================
# DISPATCH
# ============================================================


LOADERS = {
    ".pdf": load_pdf,
    ".docx": load_docx,
    ".xlsx": load_xlsx,
    ".csv": load_csv,
    ".txt": load_text,
    ".md": load_text,
}


def load_document(
    path: Path,
) -> list[LoadedBlock]:
    """
    Dispatch a file to the appropriate loader.
    """

    loader = LOADERS.get(
        path.suffix.lower()
    )

    if loader is None:
        raise ValueError(
            f"Unsupported file type: {path.suffix}"
        )

    return loader(path)